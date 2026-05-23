#!/usr/bin/env python3
"""
load_record.py -- Load an approved lease record (JSON) into Database.xlsx.

This is the deterministic "plumbing" half of the skill. The agent does the
judgment (reading the lease, filling the JSON to match the schema). This
script then:
  * validates every enum / date / required field,
  * de-duplicates Tenants, Landlords, and Properties (reuses an existing row
    when the same entity is already in the workbook),
  * auto-increments primary keys and wires up foreign keys,
  * appends rows BELOW the 4 header rows of each sheet,
  * saves the workbook in place.

Run --dry-run first to preview the row assignments without writing.

Usage:
    python load_record.py --db "/path/Database.xlsx" --record record.json
    python load_record.py --db "/path/Database.xlsx" --record record.json --dry-run

Record JSON shape: see reference/example_record.json.
Requires openpyxl (pip install openpyxl --break-system-packages).
"""
import argparse
import json
import os
import re
import shutil
import sys
import tempfile
import time
from datetime import datetime

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl --break-system-packages")

# ---- Mount-robust I/O -----------------------------------------------------
# Some environments expose Database.xlsx over a virtualized mount that drops
# in and out (the file lists in a directory but a read/write intermittently
# fails). These helpers copy the file to local disk, edit that copy, then copy
# it back with retries and a read-back verification. Loading the existing
# workbook (rather than rebuilding it) preserves all styling — fills, fonts,
# column widths, frozen panes — automatically.
IO_ATTEMPTS = 12
IO_DELAY = 4  # seconds between attempts


def load_workbook_robust(path):
    """Copy `path` to a local temp file (retrying) and load it with openpyxl."""
    last = None
    fd, tmp = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    for i in range(1, IO_ATTEMPTS + 1):
        try:
            shutil.copy(path, tmp)
            return openpyxl.load_workbook(tmp), tmp
        except Exception as e:  # FileNotFoundError, BadZipFile, etc.
            last = e
            print(f"  read attempt {i}/{IO_ATTEMPTS} failed ({type(e).__name__}); retrying…")
            time.sleep(IO_DELAY)
    sys.exit(f"Could not read {path} after {IO_ATTEMPTS} attempts: {last}")


def save_workbook_robust(wb, path):
    """Save `wb` locally, copy it back to `path` (retrying), and verify."""
    fd, tmp = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    wb.save(tmp)
    last = None
    for i in range(1, IO_ATTEMPTS + 1):
        try:
            shutil.copy(tmp, path)
            openpyxl.load_workbook(path)  # read-back verify it landed intact
            os.remove(tmp)
            return
        except Exception as e:
            last = e
            print(f"  write attempt {i}/{IO_ATTEMPTS} failed ({type(e).__name__}); retrying…")
            time.sleep(IO_DELAY)
    sys.exit(f"Could not write {path} after {IO_ATTEMPTS} attempts "
             f"(local copy kept at {tmp}): {last}")

# ---- Database.xlsx layout -------------------------------------------------
# Each sheet has 4 header rows: 1=column names, 2=PK/FK, 3=type, 4=description.
# Real data begins on row 5.
HEADER_ROWS = 4

# ---- Allowed enum values (from Completed - Schema.xlsx) --------------------
ENUMS = {
    "property_type": {"OFF", "MED", "IND", "RET", "COW", "LAB", "DC", "GND"},
    "size_unit": {"RSF", "ACRE"},
    "lease_type": {"NNN", "GRS"},
    "base_rent_frequency": {"ANNUAL", "MONTHLY"},
    "security_deposit_type": {"Cash", "CL", "NONE"},
    "renewal_rent_basis": {"FMR", "FIX", "CPI", "NONE"},
}

# Columns expected in each sheet (header row 1), used to map field -> column.
SHEET_COLS = {
    "Tenants": ["tenant_id", "tenant_name", "tenant_entity_type",
                "tenant_contact_name", "tenant_contact_email",
                "tenant_contact_phone", "notes"],
    "Landlords": ["landlord_id", "landlord_name", "landlord_entity_type",
                  "landlord_contact_name", "landlord_contact_email",
                  "landlord_contact_phone", "notes"],
    "Properties": ["property_id", "premises_address", "property_type",
                   "size", "size_unit", "notes"],
    "Leases": ["lease_id", "tenant_id", "landlord_id", "property_id",
               "commencement_date", "expiration_date", "term_months",
               "lease_type", "base_rent", "base_rent_frequency",
               "escalation_pct", "security_deposit_amount",
               "security_deposit_type", "permitted_use", "notes"],
    "RenewalOptions": ["option_id", "lease_id", "option_sequence",
                       "renewal_term_months", "renewal_rent_basis",
                       "renewal_rent_value", "notice_min_months",
                       "notice_max_months", "notes"],
}


def fail(msg):
    sys.exit(f"VALIDATION ERROR: {msg}")


def norm(s):
    """Normalize a name/address for dedup comparison."""
    if s is None:
        return ""
    return re.sub(r"\s+", " ", str(s)).strip().lower().rstrip(".")


def check_enum(field, value):
    if value is None:
        return
    if value not in ENUMS[field]:
        fail(f"{field}={value!r} not in allowed {sorted(ENUMS[field])}")


def check_date(field, value):
    if value in (None, ""):
        return
    try:
        datetime.strptime(str(value), "%Y-%m-%d")
    except ValueError:
        fail(f"{field}={value!r} is not ISO 8601 (YYYY-MM-DD)")


def header_map(ws):
    """Map column name -> 1-based column index from header row 1."""
    m = {}
    for col in range(1, ws.max_column + 1):
        name = ws.cell(row=1, column=col).value
        if name:
            m[str(name).strip()] = col
    return m


# Tracks rows virtually added per sheet during --dry-run so the preview shows
# accurate sequential PKs / row numbers when several rows go to one sheet.
_DRY_LEDGER = {}


def data_rows(ws, pk_col):
    """Yield (row_idx, pk_value) for rows that hold real data (PK present)."""
    for r in range(HEADER_ROWS + 1, ws.max_row + 1):
        v = ws.cell(row=r, column=pk_col).value
        if v not in (None, ""):
            yield r, v


def next_pk(ws, pk_col):
    mx = 0
    for _, v in data_rows(ws, pk_col):
        try:
            mx = max(mx, int(v))
        except (TypeError, ValueError):
            pass
    return mx + _DRY_LEDGER.get(ws.title, 0) + 1


def first_empty_row(ws, pk_col):
    last = HEADER_ROWS
    for r, _ in data_rows(ws, pk_col):
        last = r
    return last + _DRY_LEDGER.get(ws.title, 0) + 1


def find_existing(ws, hmap, match_field, match_value, id_field):
    """Return existing PK if a row already matches on match_field, else None."""
    mcol = hmap[match_field]
    idcol = hmap[id_field]
    target = norm(match_value)
    for r, _ in data_rows(ws, idcol):
        if norm(ws.cell(row=r, column=mcol).value) == target and target != "":
            return int(ws.cell(row=r, column=idcol).value)
    return None


def write_row(ws, hmap, pk_col, values: dict, dry, label):
    """Append a row mapping field-name -> value. Returns assigned PK."""
    pk_name = SHEET_COLS[ws.title][0]
    pk = next_pk(ws, pk_col)
    values = dict(values)
    values[pk_name] = pk
    row = first_empty_row(ws, pk_col)
    preview = {k: values.get(k) for k in SHEET_COLS[ws.title]}
    print(f"  [{label}] row {row}: {preview}")
    if dry:
        _DRY_LEDGER[ws.title] = _DRY_LEDGER.get(ws.title, 0) + 1
    else:
        for field, col in hmap.items():
            if field in values:
                ws.cell(row=row, column=col, value=values[field])
    return pk


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--db", required=True, help="Path to Database.xlsx")
    ap.add_argument("--record", required=True, help="Path to approved record JSON")
    ap.add_argument("--dry-run", action="store_true",
                    help="Validate and preview without writing")
    args = ap.parse_args()

    with open(args.record, encoding="utf-8") as f:
        rec = json.load(f)

    t = rec["tenant"]
    l = rec["landlord"]
    p = rec["property"]
    ls = rec["lease"]
    ros = rec.get("renewal_options", []) or []

    # ---- validation -------------------------------------------------------
    if not t.get("tenant_name"):
        fail("tenant.tenant_name is required")
    if not l.get("landlord_name"):
        fail("landlord.landlord_name is required")
    if not p.get("premises_address"):
        fail("property.premises_address is required")
    check_enum("property_type", p.get("property_type"))
    check_enum("size_unit", p.get("size_unit"))
    check_enum("lease_type", ls.get("lease_type"))
    check_enum("base_rent_frequency", ls.get("base_rent_frequency"))
    check_enum("security_deposit_type", ls.get("security_deposit_type"))
    check_date("commencement_date", ls.get("commencement_date"))
    check_date("expiration_date", ls.get("expiration_date"))
    for ro in ros:
        check_enum("renewal_rent_basis", ro.get("renewal_rent_basis"))

    # Soft guardrail: escalation is the percent value itself (2.5 = 2.5%).
    # A value in (0,1) is almost certainly a fraction mistake (0.025).
    esc = ls.get("escalation_pct")
    if isinstance(esc, (int, float)) and 0 < esc < 1:
        print(f"WARNING: escalation_pct={esc} looks like a fraction. Record the "
              f"percent value itself, e.g. 2.5 for 2.5% (see schema_mapping.md).")
    print("Validation passed.")

    wb, _tmp = load_workbook_robust(args.db)
    for s in SHEET_COLS:
        if s not in wb.sheetnames:
            fail(f"Database.xlsx is missing sheet {s!r}")

    print(f"\n{'DRY RUN -- no changes will be saved' if args.dry_run else 'LOADING'}:")

    # ---- Tenants (dedup on tenant_name) -----------------------------------
    ws = wb["Tenants"]; hm = header_map(ws); pk = hm["tenant_id"]
    tid = find_existing(ws, hm, "tenant_name", t["tenant_name"], "tenant_id")
    if tid:
        print(f"  [Tenant] reuse existing tenant_id={tid} ({t['tenant_name']})")
    else:
        tid = write_row(ws, hm, pk, t, args.dry_run, "Tenant")

    # ---- Landlords (dedup on landlord_name) -------------------------------
    ws = wb["Landlords"]; hm = header_map(ws); pk = hm["landlord_id"]
    lid = find_existing(ws, hm, "landlord_name", l["landlord_name"], "landlord_id")
    if lid:
        print(f"  [Landlord] reuse existing landlord_id={lid} ({l['landlord_name']})")
    else:
        lid = write_row(ws, hm, pk, l, args.dry_run, "Landlord")

    # ---- Properties (dedup on premises_address) ---------------------------
    ws = wb["Properties"]; hm = header_map(ws); pk = hm["property_id"]
    pid = find_existing(ws, hm, "premises_address", p["premises_address"], "property_id")
    if pid:
        print(f"  [Property] reuse existing property_id={pid}")
    else:
        pid = write_row(ws, hm, pk, p, args.dry_run, "Property")

    # ---- Leases (always new) ----------------------------------------------
    ws = wb["Leases"]; hm = header_map(ws); pk = hm["lease_id"]
    lease_vals = dict(ls)
    lease_vals.update({"tenant_id": tid, "landlord_id": lid, "property_id": pid})
    lease_id = write_row(ws, hm, pk, lease_vals, args.dry_run, "Lease")

    # ---- RenewalOptions (one row per option) ------------------------------
    ws = wb["RenewalOptions"]; hm = header_map(ws); pk = hm["option_id"]
    for ro in sorted(ros, key=lambda x: x.get("option_sequence", 0)):
        ro_vals = dict(ro)
        ro_vals["lease_id"] = lease_id
        write_row(ws, hm, pk, ro_vals, args.dry_run, "RenewalOption")
    if not ros:
        print("  [RenewalOption] none (lease has no renewal option)")

    if args.dry_run:
        print("\nDry run complete. Re-run without --dry-run to write.")
    else:
        save_workbook_robust(wb, args.db)
        print(f"\nSaved -> {args.db}")
        print(f"Lease loaded as lease_id={lease_id} "
              f"(tenant_id={tid}, landlord_id={lid}, property_id={pid}).")


if __name__ == "__main__":
    main()
