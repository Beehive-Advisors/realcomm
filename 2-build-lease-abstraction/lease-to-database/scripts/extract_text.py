#!/usr/bin/env python3
"""
extract_text.py -- Dump a lease .docx to plain text for abstraction.

Emits every paragraph in order, then every table (including the Article 1
"Key Defined Terms" grid and the Base Rent schedule) rendered as
pipe-delimited rows. The agent reads this text to populate the lease record.

Usage:
    python extract_text.py "/path/to/Some_Lease.docx"
    python extract_text.py "/path/to/Some_Lease.docx" --out lease.txt

Notes:
- Tables matter: in these leases the structured facts (parties, dates, size,
  base-rent schedule, security deposit, renewal option) live in tables, not
  prose. Always read the table dump, not just the paragraphs.
- Requires python-docx (pip install python-docx --break-system-packages).
"""
import argparse
import sys

try:
    import docx
except ImportError:
    sys.exit("python-docx is required: pip install python-docx --break-system-packages")


def extract(path: str) -> str:
    d = docx.Document(path)
    out = []

    out.append("==================== PARAGRAPHS ====================")
    for p in d.paragraphs:
        if p.text.strip():
            out.append(p.text.strip())

    for ti, t in enumerate(d.tables):
        out.append("")
        out.append(f"==================== TABLE {ti} ====================")
        for r in t.rows:
            cells = [c.text.strip() for c in r.cells]
            if any(cells):
                out.append(" | ".join(cells))

    return "\n".join(out)


def main():
    ap = argparse.ArgumentParser(description="Extract text + tables from a lease .docx")
    ap.add_argument("docx_path", help="Path to the lease .docx")
    ap.add_argument("--out", help="Optional output .txt path; prints to stdout if omitted")
    args = ap.parse_args()

    text = extract(args.docx_path)

    if args.out:
        with open(args.out, "w", encoding="utf-8") as f:
            f.write(text)
        print(f"Wrote {len(text):,} chars to {args.out}")
    else:
        print(text)


if __name__ == "__main__":
    main()
